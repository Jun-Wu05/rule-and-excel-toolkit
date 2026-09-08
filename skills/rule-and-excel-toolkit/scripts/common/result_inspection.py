from __future__ import annotations

import base64
import json
from pathlib import Path
from typing import Any


def _decode_rule_file(path: str) -> list[dict[str, Any]]:
    raw = Path(path).read_text(encoding="utf-8")
    clean = "".join(raw.split())
    decoded = base64.b64decode(clean).decode("utf-8")
    data = json.loads(decoded)
    if not isinstance(data, list):
        raise ValueError("decoded rule payload is not a JSON array")
    return [item for item in data if isinstance(item, dict)]


def _collect_ref_ids(value: Any) -> set[str]:
    refs: set[str] = set()
    if isinstance(value, dict):
        for key, item in value.items():
            if key == "ref" and isinstance(item, list):
                refs.update(v for v in item if isinstance(v, str) and v)
            else:
                refs.update(_collect_ref_ids(item))
    elif isinstance(value, list):
        for item in value:
            refs.update(_collect_ref_ids(item))
    return refs


def _collect_info_ids(rules: list[dict[str, Any]]) -> set[str]:
    refs: set[str] = set()
    for rule in rules:
        for section in ("normalize", "properties"):
            for entry in rule.get(section) or []:
                if not isinstance(entry, dict):
                    continue
                for info in entry.get("info") or []:
                    if isinstance(info, dict) and isinstance(info.get("id"), str) and info.get("id"):
                        refs.add(info["id"])
    return refs


def inspect_rule_file(path: str) -> dict[str, Any]:
    rules = _decode_rule_file(path)
    ids = [r.get("id") for r in rules if isinstance(r.get("id"), str) and r.get("id")]
    known = set(ids)
    refs = _collect_ref_ids(rules) | _collect_info_ids(rules)
    unknown = sorted(refs - known)
    return {
        "rule_count": len(rules),
        "top_level_id_count": len(ids),
        "duplicate_top_level_ids": len(ids) - len(set(ids)),
        "unknown_refs": unknown,
        "unknown_ref_count": len(unknown),
        "top_level_ids": ids,
    }


def inspect_rule_transition(input_path: str, output_path: str) -> tuple[dict[str, Any], dict[str, Any]]:
    before = inspect_rule_file(input_path)
    after = inspect_rule_file(output_path)
    before_unknown = set(before["unknown_refs"])
    after_unknown = set(after["unknown_refs"])
    new_unknown = sorted(after_unknown - before_unknown)
    before_ids = before["top_level_ids"]
    after_ids = after["top_level_ids"]
    changed_ids = sum(1 for old, new in zip(before_ids, after_ids) if old != new)
    stats = {
        "input_rule_count": before["rule_count"],
        "output_rule_count": after["rule_count"],
        "duplicate_top_level_ids": after["duplicate_top_level_ids"],
        "unknown_refs_before": before["unknown_ref_count"],
        "unknown_refs_after": after["unknown_ref_count"],
        "new_unknown_refs": len(new_unknown),
        "changed_top_level_ids": changed_ids,
    }
    checks = {
        "output_decodable": True,
        "top_level_ids_unique": after["duplicate_top_level_ids"] == 0,
        "new_unknown_refs_empty": len(new_unknown) == 0,
        "new_unknown_refs": new_unknown,
    }
    return stats, checks


def inspect_excel(input_path: str | None, output_path: str) -> tuple[dict[str, Any], dict[str, Any]]:
    from openpyxl import load_workbook

    wb = load_workbook(output_path, read_only=True, data_only=True)
    sheet_rows: dict[str, int] = {}
    sheet_columns: dict[str, list[str]] = {}
    for ws in wb.worksheets:
        rows = max(ws.max_row - 1, 0)
        sheet_rows[ws.title] = rows
        header = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ())
        sheet_columns[ws.title] = [str(v) for v in header if v is not None]
    stats: dict[str, Any] = {
        "sheet_count": len(wb.sheetnames),
        "sheet_names": list(wb.sheetnames),
        "sheet_rows": sheet_rows,
        "sheet_columns": sheet_columns,
    }
    if input_path:
        in_wb = load_workbook(input_path, read_only=True, data_only=True)
        stats["input_rows"] = max(in_wb.active.max_row - 1, 0)
    checks = {
        "output_readable": True,
        "has_sheet": len(wb.sheetnames) > 0,
    }
    return stats, checks


def checks_pass(checks: dict[str, Any]) -> bool:
    for key, value in checks.items():
        if key.endswith("_refs") or key == "new_unknown_refs":
            continue
        if isinstance(value, bool) and not value:
            return False
    return True
