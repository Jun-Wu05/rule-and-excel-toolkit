#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""按指定字段拆 Sheet 的日志加工脚本。

默认保持旧行为：保留全部源列，写“原始全量数据”+ 去重页 + 每值全量页。
新增：
- --keep-columns: 仅保留指定源列，再追加提取字段；
- --no-full-sheet: 不写“原始全量数据”Sheet；
- --tail-fields: 指定 plain KV 中需要一直取到日志末尾的字段，默认 raw_data；
- 支持 JSON、带引号 KV、不带引号 KV。
"""

import argparse
import re
import sys
from pathlib import Path

import pandas as pd
from openpyxl import Workbook

SCRIPT_DIR = Path(__file__).resolve().parent
if str(SCRIPT_DIR) not in sys.path:
    sys.path.insert(0, str(SCRIPT_DIR))

from common.log_fields import extract_fields

_ILLEGAL_CHARS_RE = re.compile(r'[\x00-\x08\x0b\x0c\x0e-\x1f]')


def clean_for_excel(value):
    if isinstance(value, str):
        return _ILLEGAL_CHARS_RE.sub('', value)
    return value


def safe_sheet_name(name, used):
    base = str(name) if name is not None and str(name) != "" else "空"
    base = re.sub(r'[\:\\/?*\[\]]', '_', base).strip()[:31]
    candidate = base
    i = 2
    while candidate in used:
        suffix = f"_{i}"
        candidate = base[:31 - len(suffix)] + suffix
        i += 1
    used.add(candidate)
    return candidate


def _parse_csv(raw):
    return [x.strip() for x in (raw or "").split(",") if x.strip()]


def process(
    input_path,
    output_path,
    log_column,
    target_fields,
    split_field="deviceAddress",
    keep_columns=None,
    include_full_sheet=True,
    tail_fields=None,
):
    print(f"[INFO] INPUT={input_path}")
    df = pd.read_excel(input_path, dtype=object)
    if log_column not in df.columns:
        raise ValueError(f"未找到列 '{log_column}'，可用列: {list(df.columns)}")

    extracted = df[log_column].apply(lambda s: extract_fields(s, target_fields, tail_fields=tail_fields))
    field_df = pd.DataFrame(extracted.tolist(), index=df.index)

    if keep_columns is None:
        base_df = df.copy()
    else:
        missing = [c for c in keep_columns if c not in df.columns]
        if missing:
            raise ValueError(f"--keep-columns 包含不存在的源列: {missing}")
        base_df = df[keep_columns].copy()

    base_df = base_df.drop(columns=[c for c in target_fields if c in base_df.columns], errors="ignore")
    out = pd.concat([base_df, field_df[target_fields]], axis=1)
    out = out.applymap(clean_for_excel)

    if split_field not in out.columns:
        raise ValueError(f"未找到拆分字段 '{split_field}'，可用列: {list(out.columns)}")

    dedup = out.drop_duplicates(subset=[split_field], keep="first").reset_index(drop=True)
    distinct = [v for v in dict.fromkeys(out[split_field].fillna("").astype(str).tolist()) if v]

    print(f"[STATS] INPUT_ROWS={len(df)}")
    print(f"[STATS] OUTPUT_COLUMNS={list(out.columns)}")
    print(f"[STATS] DEDUP_ROWS={len(dedup)}")
    print(f"[STATS] UNIQUE_{split_field}={len(distinct)}")

    wb = Workbook()
    wb.remove(wb.active)
    used = set()

    def write_sheet(name, frame):
        sn = safe_sheet_name(name, used)
        ws = wb.create_sheet(title=sn)
        ws.append(list(frame.columns))
        for row in frame.itertuples(index=False, name=None):
            ws.append(["" if (v is None or (isinstance(v, float) and pd.isna(v))) else v for v in row])
        for i, col in enumerate(frame.columns, 1):
            ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = max(12, min(60, len(str(col)) + 4))
        return sn

    if include_full_sheet:
        write_sheet("原始全量数据", out)
    write_sheet(f"{split_field}去重", dedup)
    for value in distinct:
        sub = out[out[split_field].fillna("").astype(str) == value].reset_index(drop=True)
        write_sheet(value, sub)

    wb.save(output_path)
    print(f"[OUTPUT] PATH={output_path}")
    print(f"[STATS] SHEETS={len(wb.sheetnames)}")
    for field in target_fields:
        non_empty = int((out[field].fillna("").astype(str) != "").sum())
        print(f"[STATS] FIELD_{field}_NONEMPTY={non_empty}")


def main():
    p = argparse.ArgumentParser(description="提取日志字段并按指定字段拆多 Sheet")
    p.add_argument("input")
    p.add_argument("output")
    p.add_argument("--log-column", default="原始日志")
    p.add_argument("--fields", default="deviceName,productVendorName,deviceSendProductName,dvcAddress,rawEvent,deviceAddress,dataType")
    p.add_argument("--split-field", default="deviceAddress")
    p.add_argument("--keep-columns", default=None, help="仅保留这些源列，逗号分隔；不传则保留全部源列")
    p.add_argument("--tail-fields", default="raw_data", help="plain KV 中取到日志末尾的字段，逗号分隔；默认 raw_data")
    p.add_argument("--no-full-sheet", action="store_true", help="不输出原始全量数据 Sheet")
    a = p.parse_args(sys.argv[1:])

    process(
        a.input,
        a.output,
        a.log_column,
        _parse_csv(a.fields),
        a.split_field,
        keep_columns=None if a.keep_columns is None else _parse_csv(a.keep_columns),
        include_full_sheet=not a.no_full_sheet,
        tail_fields=_parse_csv(a.tail_fields),
    )


if __name__ == "__main__":
    main()
