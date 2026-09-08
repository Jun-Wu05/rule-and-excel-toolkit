from __future__ import annotations

import base64
import json
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parents[1]
SCRIPTS = ROOT / "skills" / "rule-and-excel-toolkit" / "scripts"
TOOLKIT = SCRIPTS / "toolkit.py"


def write_rule_fixture(path: Path) -> None:
    child_id = "11111111-1111-4111-8111-111111111111"
    entry_id = "22222222-2222-4222-8222-222222222222"
    rules = [
        {
            "id": child_id,
            "name": "子规则",
            "subResolver": 1,
            "normalize": [{"field": "location_country", "result": ["CN"], "info": []}],
            "properties": [],
            "parser": {"filter": []},
        },
        {
            "id": entry_id,
            "name": "入口规则",
            "subResolver": 0,
            "normalize": [{"field": "original_log", "result": ["sample"], "info": []}],
            "properties": [],
            "parser": {
                "filter": [
                    {
                        "name": "redirect",
                        "ref": [child_id],
                        "cases": [
                            {
                                "rule": {
                                    "id": "33333333-3333-4333-8333-333333333333",
                                    "ref": [child_id],
                                }
                            }
                        ],
                    }
                ]
            },
        },
    ]
    raw = json.dumps(rules, ensure_ascii=False).encode("utf-8")
    path.write_text(base64.b64encode(raw).decode("ascii"), encoding="utf-8")


def write_log_fixture(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet0"
    ws.append(["原始日志", "其他列"])
    ws.append([json.dumps({"deviceName": "dev-a", "deviceAddress": "10.0.0.1", "name": "evt-a", "log_msg": "m1", "alarmExtendFieldsStrategyName": "yes"}, ensure_ascii=False), "x"])
    ws.append([json.dumps({"deviceName": "dev-b", "deviceAddress": "10.0.0.2", "name": "evt-b", "log_msg": "m2"}, ensure_ascii=False), "y"])
    ws.append(["dev_ip=10.0.0.3;raw_data=hello;inner=1", "z"])
    wb.save(path)


def write_device_fixture(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.append(["设备IP", "设备名称", "设备厂商"])
    ws.append(["10.0.0.1", "device-1", "vendor-1"])
    wb.save(path)


def write_join_log_fixture(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.append(["设备描述", "原始保留", "原始日志"])
    ws.append(["10.0.0.1", "keep", "hello"])
    wb.save(path)


class CliE2ETests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self.tmp.name)
        self.rule = self.dir / "rules.txt"
        self.log = self.dir / "logs.xlsx"
        self.device = self.dir / "devices.xlsx"
        self.join_log = self.dir / "join-log.xlsx"
        write_rule_fixture(self.rule)
        write_log_fixture(self.log)
        write_device_fixture(self.device)
        write_join_log_fixture(self.join_log)

    def tearDown(self):
        self.tmp.cleanup()

    def run_cli(self, *args: str):
        result = subprocess.run(
            [sys.executable, str(TOOLKIT), *args, "--format", "json"],
            capture_output=True,
            text=True,
        )
        try:
            payload = json.loads(result.stdout)
        except Exception as exc:
            self.fail(f"invalid JSON output: {exc}\nstdout={result.stdout}\nstderr={result.stderr}")
        self.assertEqual(result.returncode, 0, f"payload={payload}\nstderr={result.stderr}")
        self.assertEqual(payload["status"], "success", payload)
        return payload

    def assert_output_exists(self, payload):
        output = payload.get("output")
        self.assertTrue(output, payload)
        self.assertTrue(Path(output).exists(), output)

    def test_rule_hierarchy(self):
        p = self.run_cli("rule", "hierarchy", "--input", str(self.rule), "--prefix", "T_", "--verify")
        self.assert_output_exists(p)
        self.assertEqual(p["verification"]["status"], "pass")

    def test_rule_clear_field(self):
        p = self.run_cli("rule", "clear-field", "--input", str(self.rule), "--fields", "location_country", "--verify")
        self.assert_output_exists(p)
        self.assertEqual(p["verification"]["status"], "pass")

    def test_rule_reuuid(self):
        p = self.run_cli("rule", "reuuid", "--input", str(self.rule), "--verify")
        self.assert_output_exists(p)
        self.assertEqual(p["verification"]["status"], "pass")
        self.assertEqual(p["stats"]["new_unknown_refs"], 0)

    def test_rule_clone_entry(self):
        p = self.run_cli("rule", "clone-entry", "--input", str(self.rule), "--suffixes", "_copy", "--verify")
        self.assert_output_exists(p)
        self.assertEqual(p["verification"]["status"], "pass")

    def test_rule_link(self):
        p = self.run_cli("rule", "link", "--input", str(self.rule), "--verify")
        self.assert_output_exists(p)
        self.assertEqual(p["verification"]["status"], "pass")

    def test_excel_inspect(self):
        p = self.run_cli("excel", "inspect", "--input", str(self.log), "--fields", "deviceName,deviceAddress")
        self.assertEqual(p["command"], "excel.inspect")

    def test_excel_extract(self):
        p = self.run_cli("excel", "extract", "--input", str(self.log), "--fields", "deviceName,deviceAddress,name", "--verify")
        self.assert_output_exists(p)
        self.assertEqual(p["verification"]["status"], "pass")

    def test_excel_filter(self):
        p = self.run_cli("excel", "filter", "--input", str(self.log), "--keyword", "alarmExtendFieldsStrategyName", "--verify")
        self.assert_output_exists(p)
        self.assertEqual(p["verification"]["status"], "pass")

    def test_excel_split(self):
        p = self.run_cli(
            "excel", "split", "--input", str(self.log),
            "--fields", "deviceName,deviceAddress,name",
            "--split-field", "deviceAddress",
            "--keep-columns", "原始日志",
            "--no-full-sheet",
            "--verify",
        )
        self.assert_output_exists(p)
        self.assertEqual(p["verification"]["status"], "pass")
        wb = load_workbook(p["output"], read_only=True)
        self.assertIn("deviceAddress去重", wb.sheetnames)
        wb.close()

    def test_excel_split_default_keeps_only_log_column(self):
        p = self.run_cli(
            "excel", "split", "--input", str(self.log),
            "--fields", "deviceName,deviceAddress,name",
            "--split-field", "deviceAddress",
            "--no-full-sheet",
            "--verify",
        )
        self.assert_output_exists(p)
        wb = load_workbook(p["output"], read_only=True)
        header = [c.value for c in next(wb["deviceAddress去重"].iter_rows(min_row=1, max_row=1))]
        self.assertEqual(header, ["deviceName", "deviceAddress", "name", "原始日志"])
        wb.close()

    def test_excel_split_keep_all_columns(self):
        p = self.run_cli(
            "excel", "split", "--input", str(self.log),
            "--fields", "deviceName,deviceAddress,name",
            "--split-field", "deviceAddress",
            "--keep-all-columns",
            "--no-full-sheet",
            "--verify",
        )
        self.assert_output_exists(p)
        wb = load_workbook(p["output"], read_only=True)
        header = [c.value for c in next(wb["deviceAddress去重"].iter_rows(min_row=1, max_row=1))]
        self.assertEqual(header, ["deviceName", "deviceAddress", "name", "原始日志", "其他列"])
        wb.close()

    def test_excel_dedup(self):
        p = self.run_cli("excel", "dedup", "--input", str(self.log), "--fields", "log_msg", "--dedup-cols", "log_msg", "--verify")
        self.assert_output_exists(p)
        self.assertEqual(p["verification"]["status"], "pass")

    def test_excel_join(self):
        p = self.run_cli(
            "excel", "join", "--input", str(self.device),
            "--log-files", str(self.join_log),
            "--verify",
        )
        self.assert_output_exists(p)
        self.assertEqual(p["verification"]["status"], "pass")


if __name__ == "__main__":
    unittest.main()
